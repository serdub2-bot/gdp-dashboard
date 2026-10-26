#!/usr/bin/env python3
"""Собирает Excel-книгу с зонами Третьей лиги России 2007.

Книга нужна, чтобы заводить клубы в Data Editor FM2008 построчно:
один лист на зону, плюс сводный лист.

Запуск:  python3 src/make_xlsx.py
Результат: data/third_league_2007.xlsx
"""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

HEAD = ["№", "Клуб", "Город", "Статус", "Стадион", "Вместимость", "Заведён в FM"]
WIDTHS = [6, 34, 24, 12, 28, 13, 14]

THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F6FEB")
HDR_FONT = Font(bold=True, color="FFFFFF")
OUT_FILL = PatternFill("solid", fgColor="FFF5D6")


def safe_sheet(name: str) -> str:
    for ch in '[]:*?/\\':
        name = name.replace(ch, "-")
    return name[:31]


def main() -> None:
    rows = list(csv.DictReader((DATA / "clubs_2007.csv").open(encoding="utf-8")))

    # порядок зон — как в файле
    order: list[tuple[str, str]] = []
    for r in rows:
        key = (r["zone"], r["subzone"])
        if key not in order:
            order.append(key)

    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.append(["Зона", "Подзона", "Команд"])
    for zone, sub in order:
        n = sum(1 for r in rows if (r["zone"], r["subzone"]) == (r["zone"], r["subzone"])
                and r["zone"] == zone and r["subzone"] == sub)
        ws.append([zone, sub or "—", n])
    ws.append(["ВСЕГО", "", len(rows)])
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal="center")
    for i, w in enumerate([24, 26, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws["A" + str(ws.max_row)].font = Font(bold=True)

    used: set[str] = set()
    for zone, sub in order:
        title = safe_sheet(f"{zone}" + (f" — {sub}" if sub else ""))
        base, k = title, 2
        while title in used:
            title = f"{base[:28]}_{k}"
            k += 1
        used.add(title)

        sh = wb.create_sheet(title)
        sh.append(HEAD)
        for c in sh[1]:
            c.fill, c.font = HDR_FILL, HDR_FONT
            c.alignment = Alignment(horizontal="center", vertical="center")
        for i, w in enumerate(WIDTHS, start=1):
            sh.column_dimensions[get_column_letter(i)].width = w
        sh.freeze_panes = "A2"

        for r in rows:
            if (r["zone"], r["subzone"]) != (zone, sub):
                continue
            sh.append([r["position"], r["club"], r["city"], r["status"], "", "", ""])
            row = sh[sh.max_row]
            for c in row:
                c.border = BORDER
                c.alignment = Alignment(vertical="center")
            if r["status"] != "участник":
                for c in row:
                    c.fill = OUT_FILL
            row[0].alignment = Alignment(horizontal="center", vertical="center")

    out = DATA / "third_league_2007.xlsx"
    wb.save(out)

    # отдельные CSV по зонам
    per_zone = DATA / "zones_csv"
    per_zone.mkdir(exist_ok=True)
    for f in per_zone.glob("*.csv"):
        f.unlink()
    idx = 0
    for zone, sub in order:
        idx += 1
        rows_z = [r for r in rows if (r["zone"], r["subzone"]) == (zone, sub)]
        name = f"{idx:02d}_{zone}" + (f"_{sub}" if sub else "")
        name = "".join(ch if ch.isalnum() or ch in "-_" else "-" for ch in name).replace("--", "-")
        with (per_zone / f"{name}.csv").open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
            w.writeheader()
            w.writerows(rows_z)

    print(f"{out.name}: листов {len(wb.sheetnames)}, строк {len(rows)}")
    print(f"CSV по зонам: {len(order)} файлов в {per_zone.name}/")


if __name__ == "__main__":
    main()
